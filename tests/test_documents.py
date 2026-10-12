"""PDF, DOCX and OCR'd documents published over MCP as base64 (#159)."""
import base64
import io
import json
import shutil

import pypdfium2
import pytest
from docx import Document
from docx.text.paragraph import Paragraph
from PIL import Image

from src.core import node_parsers
from src.core.context_engine import ContextEngine
from src.core.mcp_adapter import build_mcp_server
from src.core.node_parsers import DOCXNodeParser, DocumentRejected, ImageNodeParser, PDFNodeParser

needs_tesseract = pytest.mark.skipif(shutil.which("tesseract") is None, reason="tesseract not installed")

PHRASE = "Quarterly revenue forecast for Atlantis"


def _pdf() -> bytes:
    stream = b"BT /F1 14 Tf 72 700 Td (%s) Tj ET" % PHRASE.encode()
    return (
        b"%%PDF-1.4\n1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n"
        b"2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj\n"
        b"3 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 612 792]/Contents 4 0 R"
        b"/Resources<</Font<</F1<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>>>>>>>endobj\n"
        b"4 0 obj<</Length %d>>stream\n%s\nendstream endobj\n"
        b"trailer<</Root 1 0 R>>\n%%%%EOF" % (len(stream), stream)
    )


def _render() -> pypdfium2.PdfBitmap:
    return pypdfium2.PdfDocument(_pdf())[0].render(scale=200 / 72, grayscale=True)


def _image() -> bytes:
    pixels = _render().to_numpy()
    return b"P5 %d %d 255\n" % (pixels.shape[1], pixels.shape[0]) + pixels.tobytes()


def _scanned_pdf() -> bytes:
    pdf = pypdfium2.PdfDocument.new()
    page = pdf.new_page(612, 792)
    image = pypdfium2.PdfImage.new(pdf)
    image.set_bitmap(_render())
    image.set_matrix(pypdfium2.PdfMatrix().scale(612, 792))
    page.insert_obj(image)
    page.gen_content()
    out = io.BytesIO()
    pdf.save(out)
    return out.getvalue()


def _docx() -> bytes:
    doc = Document()
    doc.add_paragraph(PHRASE)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


async def _publish_and_query(db, data: bytes, data_key: str, data_format: str | None = None) -> list:
    engine = ContextEngine(db=db, similarity_threshold=0.0, max_matches=10)
    await engine.initialize()
    server, _ = build_mcp_server(engine)
    await server.call_tool("contex_publish_batch", {"project_id": "docs", "items": [
        {"data_key": data_key, "data": base64.b64encode(data).decode(), "data_format": data_format},
    ]})
    result = await server.call_tool("contex_query", {
        "project_id": "docs", "query": "Atlantis revenue forecast", "top_k": 5, "threshold": 0.0,
    })
    return json.loads(result.content[0].text)["matches"]


@pytest.mark.asyncio
@pytest.mark.parametrize("make,data_key", [(_pdf, "report.pdf"), (_docx, "report.docx")])
async def test_base64_document_is_extracted(db, make, data_key):
    matches = await _publish_and_query(db, make(), data_key)
    assert matches and "Atlantis" in json.dumps(matches[0]["data"])


@pytest.mark.asyncio
async def test_bad_batch_items_are_listed_in_failed(db):
    engine = ContextEngine(db=db, similarity_threshold=0.0, max_matches=10)
    await engine.initialize()
    server, _ = build_mcp_server(engine)
    out = await server.call_tool("contex_publish_batch", {"project_id": "docs", "items": [
        {"data_key": "r.pdf", "data": "not base64!"},
        {"data_key": "no-data"},
        {"data_key": "b.py", "data": f"def {'x' * 300}():\n    return 1\n"},
    ]})
    result = json.loads(out.content[0].text)
    assert result["published"] == 1 and [f["data_key"] for f in result["failed"]] == ["r.pdf", "no-data"]


@pytest.mark.asyncio
async def test_scanned_pages_are_skipped_without_ocr(db, monkeypatch):
    monkeypatch.setattr(node_parsers, "OCR_ENABLED", False)
    assert await _publish_and_query(db, _scanned_pdf(), "scan.pdf") == []


@needs_tesseract
@pytest.mark.asyncio
async def test_scanned_pages_are_read_with_ocr(db, monkeypatch):
    monkeypatch.setattr(node_parsers, "OCR_ENABLED", True)
    matches = await _publish_and_query(db, _scanned_pdf(), "scan.pdf")
    assert matches and "Atlantis" in json.dumps(matches[0]["data"])


@needs_tesseract
@pytest.mark.asyncio
async def test_images_are_read_with_ocr(db, monkeypatch):
    monkeypatch.setattr(node_parsers, "OCR_ENABLED", True)
    matches = await _publish_and_query(db, _image(), "scan.pgm", "image")
    assert matches and "Atlantis" in json.dumps(matches[0]["data"])


def test_oversized_documents_are_refused(monkeypatch):
    monkeypatch.setattr(node_parsers, "MAX_DOCX_UNZIPPED_BYTES", 1000)
    monkeypatch.setattr(node_parsers, "MAX_PDF_PAGES", 0)
    with pytest.raises(ValueError, match="size limit"):
        DOCXNodeParser().parse(_docx())
    with pytest.raises(ValueError, match="pages"):
        PDFNodeParser().parse(_pdf())


def test_docx_paragraph_styles_are_not_resolved(monkeypatch):
    monkeypatch.setattr(Paragraph, "style", property(lambda self: 1 / 0))
    assert DOCXNodeParser().parse(_docx()).nodes[0].content == PHRASE


def test_oversized_images_are_refused_before_ocr(monkeypatch):
    monkeypatch.setattr(node_parsers, "OCR_ENABLED", True)
    monkeypatch.setattr(node_parsers, "_ocr", None)
    png = io.BytesIO()
    Image.new("1", (node_parsers.MAX_OCR_PIXELS_PER_SIDE + 1, 1)).save(png, "PNG")
    with pytest.raises(DocumentRejected, match="pixels"):
        ImageNodeParser().parse(png.getvalue())
    tiff = io.BytesIO()
    Image.new("1", (1, 1)).save(tiff, "TIFF", save_all=True, append_images=[Image.new("1", (1, 1))])
    with pytest.raises(DocumentRejected, match="one frame"):
        ImageNodeParser().parse(tiff.getvalue())
